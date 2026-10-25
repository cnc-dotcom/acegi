import os, sys, math, csv, json, datetime, traceback
from dataclasses import dataclass, asdict
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

APP_NAME = "RM CutPlan Studio"
VERSION = "0.1.0"

@dataclass
class Part:
    name: str
    width: float
    height: float
    qty: int = 1
    thickness: str = "10"
    rotate: bool = True

def shelf_nest(parts, stock_w, stock_h, gap=5.0, allow_rotate=True):
    """Simple shelf-based rectangular nesting heuristic. Units must be consistent (mm recommended).
    Returns placements per stock sheet and an estimated rectangular-shelf utilization.
    This is a starter optimizer, not a production-grade true-shape nesting engine.
    """
    expanded=[]
    for p in parts:
        if p.width <= 0 or p.height <= 0 or p.qty < 1:
            raise ValueError(f"Invalid dimensions or quantity for {p.name}")
        for i in range(p.qty):
            expanded.append({"name":p.name, "w":float(p.width), "h":float(p.height),
                             "thickness":str(p.thickness), "rotate":bool(p.rotate), "instance":i+1})
    expanded.sort(key=lambda x: max(x["w"],x["h"])*min(x["w"],x["h"]), reverse=True)
    sheets=[]
    for item in expanded:
        placed=False
        options=[(item["w"],item["h"],False)]
        if allow_rotate and item["rotate"] and abs(item["w"]-item["h"])>1e-9:
            options.append((item["h"],item["w"],True))
        # try existing shelves then a new shelf
        for si,sheet in enumerate(sheets):
            for shelf in sheet["shelves"]:
                for w,h,rot in options:
                    if h <= shelf["height"] and shelf["x"]+w <= stock_w+1e-9:
                        sheet["placements"].append({"name":item["name"],"x":shelf["x"],"y":shelf["y"],
                            "w":w,"h":h,"rotated":rot,"thickness":item["thickness"],"instance":item["instance"]})
                        shelf["x"] += w+gap
                        placed=True; break
                if placed: break
            if placed: break
            # create new shelf at current bottom
            y=sheet["used_h"]+gap
            for w,h,rot in options:
                if w <= stock_w and y+h <= stock_h:
                    sheet["shelves"].append({"x":w+gap,"y":y,"height":h})
                    sheet["placements"].append({"name":item["name"],"x":0.0,"y":y,"w":w,"h":h,
                        "rotated":rot,"thickness":item["thickness"],"instance":item["instance"]})
                    sheet["used_h"]=y+h
                    placed=True; break
            if placed: break
        if not placed:
            sheet={"placements":[],"shelves":[],"used_h":0.0}
            for w,h,rot in options:
                if w <= stock_w and h <= stock_h:
                    sheet["shelves"].append({"x":w+gap,"y":0.0,"height":h})
                    sheet["placements"].append({"name":item["name"],"x":0.0,"y":0.0,"w":w,"h":h,
                        "rotated":rot,"thickness":item["thickness"],"instance":item["instance"]})
                    sheet["used_h"]=h
                    placed=True; break
            if not placed:
                raise ValueError(f"Part {item['name']} ({item['w']} x {item['h']}) does not fit stock {stock_w} x {stock_h}")
            sheets.append(sheet)
    part_area=sum(float(p.width)*float(p.height)*int(p.qty) for p in parts)
    stock_area=len(sheets)*stock_w*stock_h
    return sheets, (100*part_area/stock_area if stock_area else 0.0)

def optimize_lengths(cut_lengths, stock_length, kerf=3.0, end_trim=0.0):
    """First-fit decreasing 1D stock optimizer. Lengths and kerf in same units."""
    pieces=[]
    for item in cut_lengths:
        length=float(item["length"]); qty=int(item.get("qty",1))
        if length<=0 or qty<1: raise ValueError("Lengths and quantities must be positive")
        if length+end_trim>stock_length: raise ValueError(f"{item.get('name','Part')} exceeds stock length")
        pieces.extend([{"name":item.get("name","Part"),"length":length} for _ in range(qty)])
    pieces.sort(key=lambda x:x["length"], reverse=True)
    bars=[]
    for p in pieces:
        chosen=None
        for b in bars:
            extra=p["length"]+(kerf if b["pieces"] else 0)
            if b["used"]+extra <= stock_length-end_trim+1e-9:
                chosen=b; break
        if chosen is None:
            chosen={"pieces":[],"used":0.0,"stock_length":float(stock_length),"end_trim":float(end_trim)}
            bars.append(chosen)
        extra=p["length"]+(kerf if chosen["pieces"] else 0)
        chosen["pieces"].append({**p,"start":chosen["used"]+(kerf if chosen["pieces"] else 0)})
        chosen["used"] += extra
    for b in bars:
        b["scrap"]=max(0.0, b["stock_length"]-b["end_trim"]-b["used"])
    total_cut=sum(p["length"] for p in pieces)
    stock_total=len(bars)*float(stock_length)
    utilization=(100*total_cut/stock_total if stock_total else 0.0)
    return bars, utilization

def read_excel_parts(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Excel support needs openpyxl. Run install_requirements.bat.")
    wb=load_workbook(path, data_only=True)
    ws=wb.active
    rows=list(ws.iter_rows(values_only=True))
    if not rows: return []
    headers=[str(x or "").strip().lower() for x in rows[0]]
    def idx(*names):
        for n in names:
            if n in headers: return headers.index(n)
        return None
    wi=idx("width","w","width_mm"); hi=idx("height","h","length","height_mm")
    qi=idx("qty","quantity"); ni=idx("name","part","part name","mark"); ti=idx("thickness","thk","thickness_mm")
    if wi is None or hi is None:
        raise ValueError("Excel headers must include Width and Height (or W and H). Optional: Name, Qty, Thickness.")
    out=[]
    for row in rows[1:]:
        if not row or row[wi] is None or row[hi] is None: continue
        out.append(Part(str(row[ni] if ni is not None and row[ni] is not None else f"Part {len(out)+1}"),
                        float(row[wi]),float(row[hi]),int(row[qi] or 1) if qi is not None else 1,
                        str(row[ti] if ti is not None and row[ti] is not None else "10")))
    return out

def read_dxf_parts(path, thickness="10"):
    try:
        import ezdxf
    except ImportError:
        raise RuntimeError("DXF support needs ezdxf. Run install_requirements.bat.")
    doc=ezdxf.readfile(path); msp=doc.modelspace(); out=[]
    # MVP supports closed LWPOLYLINE/POLYLINE and circles approximated by bounding boxes.
    for ent in msp:
        typ=ent.dxftype()
        try:
            if typ=="LWPOLYLINE" and ent.closed:
                pts=list(ent.get_points("xy"))
            elif typ=="POLYLINE" and ent.is_closed:
                pts=[(v.dxf.location.x,v.dxf.location.y) for v in ent.vertices]
            elif typ=="CIRCLE":
                c=ent.dxf.center; r=ent.dxf.radius
                pts=[(c.x-r,c.y-r),(c.x+r,c.y+r)]
            else: continue
            if len(pts)<2: continue
            xs=[p[0] for p in pts]; ys=[p[1] for p in pts]
            w=max(xs)-min(xs); h=max(ys)-min(ys)
            if w>0 and h>0:
                out.append(Part(f"DXF-{len(out)+1}",w,h,1,thickness))
        except Exception:
            continue
    if not out:
        raise ValueError("No closed polylines or circles found. DXF units are assumed to be mm.")
    return out

def export_excel(path, report_no, company, date_text, mode, summary, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb=Workbook(); ws=wb.active; ws.title="Cutting Report"
    ws.append([company]); ws.append(["CUTTING PLAN REPORT"]); ws.append(["Report No.",report_no,"Date",date_text])
    ws.append(["Plan type",mode]); ws.append([])
    for r in summary: ws.append(r)
    ws.append([])
    if rows:
        headers=list(rows[0].keys()); ws.append(headers)
        for row in rows: ws.append([row.get(h,"") for h in headers])
    for row in ws.iter_rows():
        for c in row:
            c.alignment=Alignment(vertical="center")
            if c.row in (1,2,3,4,5,7+len(summary)):
                c.font=Font(bold=True,color="FFFFFF")
                c.fill=PatternFill("solid",fgColor="17365D")
    for col in ws.columns:
        letter=col[0].column_letter
        ws.column_dimensions[letter].width=min(28,max(12,max(len(str(c.value or "")) for c in col)+2))
    ws.freeze_panes="A9"
    wb.save(path)

def export_pdf(path, report_no, company, date_text, mode, summary, rows, logo_path=""):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
        from reportlab.lib.units import mm
    except ImportError:
        raise RuntimeError("PDF support needs reportlab. Run install_requirements.bat.")
    doc=SimpleDocTemplate(path,pagesize=landscape(A4),rightMargin=12*mm,leftMargin=12*mm,topMargin=10*mm,bottomMargin=10*mm)
    styles=getSampleStyleSheet(); story=[]
    if logo_path and os.path.exists(logo_path):
        try: story.append(Image(logo_path,width=30*mm,height=18*mm,kind="proportional")); story.append(Spacer(1,3*mm))
        except Exception: pass
    story += [Paragraph(company or "Workshop",styles["Title"]),
              Paragraph("CUTTING PLAN REPORT",styles["Heading2"]),
              Paragraph(f"Report No: {report_no} &nbsp;&nbsp;&nbsp; Date: {date_text} &nbsp;&nbsp;&nbsp; Plan: {mode}",styles["Normal"]),Spacer(1,5*mm)]
    data=[[str(x) for x in r] for r in summary]
    if data:
        t=Table(data,repeatRows=0,hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#17365D")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                               ("GRID",(0,0),(-1,-1),0.4,colors.grey),("PADDING",(0,0),(-1,-1),5)]))
        story += [t,Spacer(1,5*mm)]
    if rows:
        headers=list(rows[0].keys()); data=[headers]+[[str(r.get(h,"")) for h in headers] for r in rows]
        t=Table(data,repeatRows=1,hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#17365D")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                               ("GRID",(0,0),(-1,-1),0.3,colors.grey),("FONTSIZE",(0,0),(-1,-1),7),("PADDING",(0,0),(-1,-1),3)]))
        story.append(t)
    story.append(Spacer(1,5*mm)); story.append(Paragraph("Prepared by: ____________________ &nbsp;&nbsp;&nbsp; Checked by: ____________________ &nbsp;&nbsp;&nbsp; Approved by: ____________________",styles["Normal"]))
    doc.build(story)

class App:
    def __init__(self, root):
        self.root=root; root.title(f"{APP_NAME} {VERSION}"); root.geometry("1180x760"); root.minsize(950,620)
        self.plate_parts=[]; self.length_items=[]; self.last_mode=""; self.last_summary=[]; self.last_rows=[]
        self.company=tk.StringVar(value="RM Engineering Service")
        self.report_no=tk.StringVar(value="RMC-"+datetime.datetime.now().strftime("%Y%m%d-%H%M"))
        self.stock_w=tk.StringVar(value="3000"); self.stock_h=tk.StringVar(value="1500"); self.gap=tk.StringVar(value="5")
        self.stock_len=tk.StringVar(value="12000"); self.kerf=tk.StringVar(value="3"); self.trim=tk.StringVar(value="0")
        self.logo_path=""; self.status=tk.StringVar(value="Ready. Units: mm. Save a copy of your cutting lists.")
        self._build()
    def _build(self):
        top=ttk.Frame(self.root,padding=10); top.pack(fill="x")
        ttk.Label(top,text=APP_NAME,font=("Segoe UI",18,"bold")).pack(side="left")
        ttk.Label(top,text="Free personal-use MVP",foreground="#477").pack(side="left",padx=14)
        settings=ttk.LabelFrame(self.root,text="Report setup",padding=8); settings.pack(fill="x",padx=10,pady=4)
        ttk.Label(settings,text="Company").grid(row=0,column=0,sticky="w"); ttk.Entry(settings,textvariable=self.company,width=28).grid(row=0,column=1,padx=5)
        ttk.Label(settings,text="Report number").grid(row=0,column=2,sticky="w"); ttk.Entry(settings,textvariable=self.report_no,width=24).grid(row=0,column=3,padx=5)
        ttk.Button(settings,text="Choose logo",command=self.choose_logo).grid(row=0,column=4,padx=6)
        ttk.Button(settings,text="New report no.",command=self.new_report_no).grid(row=0,column=5,padx=6)
        self.tabs=ttk.Notebook(self.root); self.tabs.pack(fill="both",expand=True,padx=10,pady=8)
        self.plate_tab=ttk.Frame(self.tabs,padding=8); self.length_tab=ttk.Frame(self.tabs,padding=8); self.report_tab=ttk.Frame(self.tabs,padding=8)
        self.tabs.add(self.plate_tab,text="Plate Nesting"); self.tabs.add(self.length_tab,text="Structural Length"); self.tabs.add(self.report_tab,text="Reports / Export")
        self._build_plate(); self._build_length(); self._build_report()
        ttk.Label(self.root,textvariable=self.status,relief="sunken",anchor="w").pack(fill="x",side="bottom")
    def choose_logo(self):
        p=filedialog.askopenfilename(title="Select company logo",filetypes=[("Image files","*.png *.jpg *.jpeg"),("All files","*.*")])
        if p: self.logo_path=p; self.status.set("Logo selected: "+os.path.basename(p))
    def new_report_no(self): self.report_no.set("RMC-"+datetime.datetime.now().strftime("%Y%m%d-%H%M%S"))
    def _build_plate(self):
        cfg=ttk.LabelFrame(self.plate_tab,text="Stock sheet settings (mm)",padding=6); cfg.pack(fill="x")
        for i,(lab,var) in enumerate([("Stock width",self.stock_w),("Stock height",self.stock_h),("Part gap / kerf allowance",self.gap)]):
            ttk.Label(cfg,text=lab).grid(row=0,column=i*2,sticky="w",padx=3); ttk.Entry(cfg,textvariable=var,width=10).grid(row=0,column=i*2+1,padx=4)
        form=ttk.LabelFrame(self.plate_tab,text="Add rectangular part (or import Excel / DXF)",padding=6); form.pack(fill="x",pady=6)
        self.p_name=tk.StringVar(value="P-001"); self.p_w=tk.StringVar(value="500"); self.p_h=tk.StringVar(value="300"); self.p_q=tk.StringVar(value="2"); self.p_t=tk.StringVar(value="10")
        for i,(lab,var) in enumerate([("Name",self.p_name),("Width",self.p_w),("Height",self.p_h),("Qty",self.p_q),("Thickness",self.p_t)]):
            ttk.Label(form,text=lab).grid(row=0,column=i*2,sticky="w"); ttk.Entry(form,textvariable=var,width=11).grid(row=0,column=i*2+1,padx=4)
        ttk.Button(form,text="Add part",command=self.add_part).grid(row=0,column=10,padx=4)
        ttk.Button(form,text="Import Excel",command=self.import_excel).grid(row=0,column=11,padx=4)
        ttk.Button(form,text="Import DXF",command=self.import_dxf).grid(row=0,column=12,padx=4)
        cols=("name","width","height","qty","thickness")
        self.plate_tree=ttk.Treeview(self.plate_tab,columns=cols,show="headings",height=10)
        for c,w in zip(cols,(220,120,120,80,100)): self.plate_tree.heading(c,text=c.title()); self.plate_tree.column(c,width=w,anchor="center")
        self.plate_tree.pack(fill="both",expand=True)
        bar=ttk.Frame(self.plate_tab); bar.pack(fill="x",pady=6)
        ttk.Button(bar,text="Remove selected",command=self.remove_part).pack(side="left")
        ttk.Button(bar,text="Optimize plate layout",command=self.run_plate).pack(side="left",padx=6)
        ttk.Label(bar,text="MVP: rectangular nesting; DXF geometry is shown as bounding boxes.").pack(side="left",padx=8)
        self.layout_canvas=tk.Canvas(self.plate_tab,height=260,bg="white",highlightthickness=1,highlightbackground="#9aa4b2")
        self.layout_canvas.pack(fill="both",expand=False,pady=6)
        ttk.Label(self.plate_tab,text="Layout preview is schematic. Confirm dimensions and clearances before cutting.",foreground="#7a4b00").pack(anchor="w")
    def _build_length(self):
        cfg=ttk.LabelFrame(self.length_tab,text="Stock settings (mm)",padding=6); cfg.pack(fill="x")
        for i,(lab,var) in enumerate([("Stock length",self.stock_len),("Saw/plasma kerf",self.kerf),("End trim allowance",self.trim)]):
            ttk.Label(cfg,text=lab).grid(row=0,column=i*2,sticky="w",padx=3); ttk.Entry(cfg,textvariable=var,width=12).grid(row=0,column=i*2+1,padx=4)
        form=ttk.LabelFrame(self.length_tab,text="Add cut length",padding=6); form.pack(fill="x",pady=6)
        self.l_name=tk.StringVar(value="Beam-001"); self.l_len=tk.StringVar(value="2450"); self.l_q=tk.StringVar(value="4")
        for i,(lab,var) in enumerate([("Part mark",self.l_name),("Length (mm)",self.l_len),("Qty",self.l_q)]):
            ttk.Label(form,text=lab).grid(row=0,column=i*2,sticky="w"); ttk.Entry(form,textvariable=var,width=14).grid(row=0,column=i*2+1,padx=4)
        ttk.Button(form,text="Add length",command=self.add_length).grid(row=0,column=6,padx=4)
        ttk.Button(form,text="Import Excel",command=self.import_length_excel).grid(row=0,column=7,padx=4)
        cols=("name","length","qty"); self.length_tree=ttk.Treeview(self.length_tab,columns=cols,show="headings",height=13)
        for c,w in zip(cols,(260,180,100)): self.length_tree.heading(c,text=c.title()); self.length_tree.column(c,width=w,anchor="center")
        self.length_tree.pack(fill="both",expand=True)
        bar=ttk.Frame(self.length_tab); bar.pack(fill="x",pady=6)
        ttk.Button(bar,text="Remove selected",command=self.remove_length).pack(side="left")
        ttk.Button(bar,text="Optimize stock lengths",command=self.run_length).pack(side="left",padx=6)
    def _build_report(self):
        ttk.Label(self.report_tab,text="Run an optimizer first. Then export a PDF or Excel report.",font=("Segoe UI",11)).pack(anchor="w",pady=6)
        ttk.Button(self.report_tab,text="Export professional PDF report",command=self.save_pdf).pack(anchor="w",pady=5)
        ttk.Button(self.report_tab,text="Export Excel report",command=self.save_xlsx).pack(anchor="w",pady=5)
        ttk.Button(self.report_tab,text="Save input template workbook",command=self.save_template).pack(anchor="w",pady=5)
        self.preview=tk.Text(self.report_tab,height=22,wrap="word"); self.preview.pack(fill="both",expand=True,pady=8)
    def add_part(self):
        try: p=Part(self.p_name.get().strip() or f"Part {len(self.plate_parts)+1}",float(self.p_w.get()),float(self.p_h.get()),int(self.p_q.get()),self.p_t.get().strip() or "10")
        except Exception: messagebox.showerror("Input error","Enter valid width, height and quantity."); return
        if p.width<=0 or p.height<=0 or p.qty<=0: messagebox.showerror("Input error","Dimensions and quantity must be positive."); return
        self.plate_parts.append(p); self.plate_tree.insert("", "end", values=(p.name,p.width,p.height,p.qty,p.thickness)); self.status.set(f"Added {p.name}")
    def import_excel(self):
        p=filedialog.askopenfilename(filetypes=[("Excel workbook","*.xlsx")])
        if not p:return
        try:
            got=read_excel_parts(p)
            for x in got: self.plate_parts.append(x); self.plate_tree.insert("", "end", values=(x.name,x.width,x.height,x.qty,x.thickness))
            self.status.set(f"Imported {len(got)} parts from Excel")
        except Exception as e: messagebox.showerror("Import failed",str(e))
    def import_dxf(self):
        p=filedialog.askopenfilename(filetypes=[("DXF drawing","*.dxf")])
        if not p:return
        try:
            got=read_dxf_parts(p,self.p_t.get() or "10")
            for x in got: self.plate_parts.append(x); self.plate_tree.insert("", "end", values=(x.name,round(x.width,2),round(x.height,2),x.qty,x.thickness))
            self.status.set(f"Imported {len(got)} DXF bounding boxes. Verify dimensions and geometry.")
        except Exception as e: messagebox.showerror("DXF import failed",str(e))
    def remove_part(self):
        for item in self.plate_tree.selection():
            vals=self.plate_tree.item(item,"values"); self.plate_parts=[p for p in self.plate_parts if not (p.name==vals[0] and str(p.width)==str(vals[1]) and str(p.height)==str(vals[2]))]
            self.plate_tree.delete(item)
    def add_length(self):
        try: item={"name":self.l_name.get().strip() or "Part","length":float(self.l_len.get()),"qty":int(self.l_q.get())}
        except Exception: messagebox.showerror("Input error","Enter a valid length and quantity."); return
        if item["length"]<=0 or item["qty"]<=0: messagebox.showerror("Input error","Length and quantity must be positive."); return
        self.length_items.append(item); self.length_tree.insert("", "end", values=(item["name"],item["length"],item["qty"]))
    def import_length_excel(self):
        p=filedialog.askopenfilename(filetypes=[("Excel workbook","*.xlsx")])
        if not p:return
        try:
            from openpyxl import load_workbook
            ws=load_workbook(p,data_only=True).active; rows=list(ws.iter_rows(values_only=True))
            if not rows: return
            headers=[str(x or "").strip().lower() for x in rows[0]]
            li=next((headers.index(x) for x in ("length","cut length","length_mm") if x in headers),None)
            qi=next((headers.index(x) for x in ("qty","quantity") if x in headers),None)
            ni=next((headers.index(x) for x in ("name","part","mark","part mark") if x in headers),None)
            if li is None: raise ValueError("Length Excel needs a Length column; optional Name and Qty.")
            count=0
            for row in rows[1:]:
                if row and row[li] is not None:
                    item={"name":str(row[ni] if ni is not None and row[ni] else f"Part {len(self.length_items)+1}"),"length":float(row[li]),"qty":int(row[qi] or 1) if qi is not None else 1}
                    self.length_items.append(item); self.length_tree.insert("", "end", values=(item["name"],item["length"],item["qty"])); count+=1
            self.status.set(f"Imported {count} length rows")
        except Exception as e: messagebox.showerror("Import failed",str(e))
    def remove_length(self):
        for item in self.length_tree.selection():
            vals=self.length_tree.item(item,"values")
            for i,x in enumerate(self.length_items):
                if x["name"]==vals[0] and str(x["length"])==str(vals[1]) and str(x["qty"])==str(vals[2]):
                    self.length_items.pop(i); break
            self.length_tree.delete(item)
    def show_result(self, mode, summary, rows):
        self.last_mode=mode; self.last_summary=summary; self.last_rows=rows
        self.preview.delete("1.0","end")
        self.preview.insert("end",f"{self.company.get()}\\nCUTTING PLAN REPORT\\nReport No: {self.report_no.get()}\\nDate: {datetime.date.today()}\\nPlan: {mode}\\n\\n")
        for r in summary: self.preview.insert("end"," | ".join(map(str,r))+"\\n")
        self.preview.insert("end","\\nDETAILS\\n")
        if rows:
            self.preview.insert("end"," | ".join(rows[0].keys())+"\\n")
            for r in rows: self.preview.insert("end"," | ".join(str(v) for v in r.values())+"\\n")
        self.tabs.select(self.report_tab); self.status.set("Optimization complete. Review report before fabrication.")
    def draw_plate_layout(self, sheets):
        if not hasattr(self, "layout_canvas"): return
        c=self.layout_canvas; c.delete("all")
        cw=max(c.winfo_width(),500); ch=max(c.winfo_height(),220)
        if not sheets:
            c.create_text(cw/2,ch/2,text="Run plate optimization to preview layout",fill="#667085")
            return
        sw=float(self.stock_w.get()); sh=float(self.stock_h.get())
        margin=24; scale=min((cw-2*margin)/sw,(ch-2*margin)/sh)
        ox=(cw-sw*scale)/2; oy=(ch-sh*scale)/2
        c.create_rectangle(ox,oy,ox+sw*scale,oy+sh*scale,outline="#17365D",width=2)
        c.create_text(ox,oy-10,anchor="w",text=f"Sheet 1 of {len(sheets)}  |  {sw:g} x {sh:g} mm",fill="#17365D")
        palette=["#cfe8ff","#d7f4e2","#ffe4c7","#e6dcff","#ffdce3","#d9f1f0"]
        for i,p in enumerate(sheets[0]["placements"]):
            x=ox+p["x"]*scale; y=oy+p["y"]*scale
            x2=x+p["w"]*scale; y2=y+p["h"]*scale
            c.create_rectangle(x,y,x2,y2,fill=palette[i%len(palette)],outline="#344054",width=1)
            if (x2-x)>35 and (y2-y)>18:
                c.create_text((x+x2)/2,(y+y2)/2,text=f'{p["name"]}\\n{p["w"]:g}x{p["h"]:g}',font=("Segoe UI",8),fill="#101828")
    def run_plate(self):
        try:
            if not self.plate_parts: raise ValueError("Add or import at least one part.")
            sheets,util=shelf_nest(self.plate_parts,float(self.stock_w.get()),float(self.stock_h.get()),float(self.gap.get()))
            rows=[]
            for i,s in enumerate(sheets,1):
                for p in s["placements"]:
                    rows.append({"Sheet":i,"Part":p["name"],"X":round(p["x"],2),"Y":round(p["y"],2),"Width":round(p["w"],2),"Height":round(p["h"],2),"Rotated":p["rotated"],"Thickness":p["thickness"]})
            summary=[["Metric","Value"],["Stock size (mm)",f"{self.stock_w.get()} x {self.stock_h.get()}"],["Sheets required",len(sheets)],["Part count",sum(p.qty for p in self.plate_parts)],["Area utilization (%)",round(util,2)],["Unutilized area (%)",round(100-util,2)],["Gap (mm)",self.gap.get()]]
            self.draw_plate_layout(sheets)
            self.show_result("Plate nesting",summary,rows)
        except Exception as e: messagebox.showerror("Optimization failed",str(e))
    def run_length(self):
        try:
            if not self.length_items: raise ValueError("Add or import at least one cut length.")
            bars,util=optimize_lengths(self.length_items,float(self.stock_len.get()),float(self.kerf.get()),float(self.trim.get()))
            rows=[]
            for i,b in enumerate(bars,1):
                for p in b["pieces"]:
                    rows.append({"Stock bar":i,"Part":p["name"],"Cut length (mm)":p["length"],"Start (mm)":round(p["start"],2),"Stock length (mm)":b["stock_length"],"Estimated scrap (mm)":round(b["scrap"],2)})
            summary=[["Metric","Value"],["Stock length (mm)",self.stock_len.get()],["Bars required",len(bars)],["Cut pieces",sum(x["qty"] for x in self.length_items)],["Cut-length utilization (%)",round(util,2)],["Kerf (mm)",self.kerf.get()],["End trim (mm)",self.trim.get()]]
            self.show_result("Structural length",summary,rows)
        except Exception as e: messagebox.showerror("Optimization failed",str(e))
    def save_pdf(self):
        if not self.last_rows: messagebox.showinfo("No report","Run an optimizer first."); return
        p=filedialog.asksaveasfilename(defaultextension=".pdf",initialfile=self.report_no.get()+".pdf",filetypes=[("PDF report","*.pdf")])
        if not p:return
        try: export_pdf(p,self.report_no.get(),self.company.get(),str(datetime.date.today()),self.last_mode,self.last_summary,self.last_rows,self.logo_path); self.status.set("PDF exported: "+p); messagebox.showinfo("Export complete","PDF report saved.")
        except Exception as e: messagebox.showerror("PDF export failed",str(e))
    def save_xlsx(self):
        if not self.last_rows: messagebox.showinfo("No report","Run an optimizer first."); return
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",initialfile=self.report_no.get()+".xlsx",filetypes=[("Excel report","*.xlsx")])
        if not p:return
        try: export_excel(p,self.report_no.get(),self.company.get(),str(datetime.date.today()),self.last_mode,self.last_summary,self.last_rows); self.status.set("Excel exported: "+p); messagebox.showinfo("Export complete","Excel report saved.")
        except Exception as e: messagebox.showerror("Excel export failed",str(e))
    def save_template(self):
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",initialfile="CutPlan_Input_Template.xlsx",filetypes=[("Excel workbook","*.xlsx")])
        if not p:return
        try:
            from openpyxl import Workbook
            wb=Workbook(); ws=wb.active; ws.title="Plate Parts"; ws.append(["Name","Width","Height","Qty","Thickness"])
            ws.append(["P-001",500,300,2,10]); wl=wb.create_sheet("Length Parts"); wl.append(["Name","Length","Qty"]); wl.append(["Beam-001",2450,4]); wb.save(p)
            messagebox.showinfo("Template saved","Template workbook saved. Use Plate Parts for plate nesting and Length Parts for structural cutting.")
        except Exception as e: messagebox.showerror("Template failed",str(e))
def main():
    root=tk.Tk(); App(root); root.mainloop()
if __name__=="__main__": main()
